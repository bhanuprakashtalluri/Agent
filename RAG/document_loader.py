"""
Universal Document Loader
Supports: PDF, Excel, CSV, Word, TXT, Markdown
"""

import os
from typing import List
from pathlib import Path
from langchain_core.documents import Document

# PDF loaders
from langchain_community.document_loaders import PyPDFLoader, UnstructuredPDFLoader

# Excel/CSV loaders
import pandas as pd

# Word loader
try:
    from docx import Document as DocxDocument
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("⚠ python-docx not installed. Word documents will not be supported.")


class DocumentLoader:
    """Universal document loader supporting multiple file types"""
    
    SUPPORTED_EXTENSIONS = {
        'pdf': 'PDF Document',
        'xlsx': 'Excel Spreadsheet',
        'xls': 'Excel Spreadsheet (Legacy)',
        'csv': 'CSV File',
        'docx': 'Word Document',
        'txt': 'Text File',
        'md': 'Markdown File',
        'json': 'JSON File'
    }
    
    def __init__(self):
        self.documents_folder = Path("RAG/documents")
        self._create_folders()
    
    def _create_folders(self):
        """Create necessary folders for document storage"""
        folders = [
            self.documents_folder / "pdfs",
            self.documents_folder / "excel",
            self.documents_folder / "word",
            self.documents_folder / "others"
        ]
        for folder in folders:
            folder.mkdir(parents=True, exist_ok=True)
    
    def load_document(self, file_path: str) -> List[Document]:
        """
        Load a document and return LangChain Document objects
        
        Args:
            file_path: Path to the document file
            
        Returns:
            List of Document objects with page_content and metadata
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = path.suffix.lower().lstrip('.')
        
        if extension not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: .{extension}\n"
                f"Supported types: {', '.join(self.SUPPORTED_EXTENSIONS.keys())}"
            )
        
        # Route to appropriate loader
        if extension == 'pdf':
            return self._load_pdf(file_path)
        elif extension in ['xlsx', 'xls']:
            return self._load_excel(file_path)
        elif extension == 'csv':
            return self._load_csv(file_path)
        elif extension == 'docx':
            return self._load_word(file_path)
        elif extension in ['txt', 'md']:
            return self._load_text(file_path)
        elif extension == 'json':
            return self._load_json(file_path)
        else:
            raise ValueError(f"Handler not implemented for .{extension}")
    
    def _load_pdf(self, file_path: str) -> List[Document]:
        """Load PDF document"""
        try:
            # Try PyPDFLoader first (faster)
            loader = PyPDFLoader(file_path)
            documents = loader.load()
            
            # Add source metadata
            for doc in documents:
                doc.metadata['source'] = file_path
                doc.metadata['type'] = 'pdf'
            
            return documents
        except Exception as e:
            print(f"PyPDFLoader failed, trying UnstructuredPDFLoader: {e}")
            # Fallback to UnstructuredPDFLoader
            try:
                loader = UnstructuredPDFLoader(file_path)
                documents = loader.load()
                for doc in documents:
                    doc.metadata['source'] = file_path
                    doc.metadata['type'] = 'pdf'
                return documents
            except Exception as e2:
                raise Exception(f"Failed to load PDF: {e2}")
    
    def _load_excel(self, file_path: str) -> List[Document]:
        """Load Excel file - converts each sheet to a document"""
        documents = []
        
        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                
                # Convert to string representation
                content = f"Sheet: {sheet_name}\n\n"
                content += df.to_string(index=False)
                
                # Also add summary statistics
                content += f"\n\nSummary:\n"
                content += f"Rows: {len(df)}\n"
                content += f"Columns: {', '.join(df.columns.tolist())}\n"
                
                doc = Document(
                    page_content=content,
                    metadata={
                        'source': file_path,
                        'sheet_name': sheet_name,
                        'type': 'excel',
                        'rows': len(df),
                        'columns': len(df.columns)
                    }
                )
                documents.append(doc)
            
            return documents
        
        except Exception as e:
            raise Exception(f"Failed to load Excel file: {e}")
    
    def _load_csv(self, file_path: str) -> List[Document]:
        """Load CSV file"""
        try:
            df = pd.read_csv(file_path)
            
            # Convert to string representation
            content = df.to_string(index=False)
            
            # Add summary
            content += f"\n\nSummary:\n"
            content += f"Rows: {len(df)}\n"
            content += f"Columns: {', '.join(df.columns.tolist())}\n"
            
            doc = Document(
                page_content=content,
                metadata={
                    'source': file_path,
                    'type': 'csv',
                    'rows': len(df),
                    'columns': len(df.columns)
                }
            )
            
            return [doc]
        
        except Exception as e:
            raise Exception(f"Failed to load CSV file: {e}")
    
    def _load_word(self, file_path: str) -> List[Document]:
        """Load Word document"""
        if not DOCX_AVAILABLE:
            raise Exception("python-docx not installed. Install with: pip install python-docx")
        
        try:
            doc = DocxDocument(file_path)
            
            # Extract all paragraphs
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            
            content = "\n\n".join(full_text)
            
            document = Document(
                page_content=content,
                metadata={
                    'source': file_path,
                    'type': 'docx',
                    'paragraphs': len(doc.paragraphs)
                }
            )
            
            return [document]
        
        except Exception as e:
            raise Exception(f"Failed to load Word document: {e}")
    
    def _load_text(self, file_path: str) -> List[Document]:
        """Load plain text or markdown file"""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            doc = Document(
                page_content=content,
                metadata={
                    'source': file_path,
                    'type': Path(file_path).suffix.lstrip('.'),
                }
            )
            
            return [doc]
        
        except Exception as e:
            raise Exception(f"Failed to load text file: {e}")
    
    def _load_json(self, file_path: str) -> List[Document]:
        """Load JSON file"""
        import json
        
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            # Convert JSON to readable string
            content = json.dumps(data, indent=2)
            
            doc = Document(
                page_content=content,
                metadata={
                    'source': file_path,
                    'type': 'json',
                }
            )
            
            return [doc]
        
        except Exception as e:
            raise Exception(f"Failed to load JSON file: {e}")
    
    def load_directory(self, directory_path: str, recursive: bool = True) -> List[Document]:
        """
        Load all supported documents from a directory
        
        Args:
            directory_path: Path to directory
            recursive: Whether to search subdirectories
            
        Returns:
            List of all loaded documents
        """
        path = Path(directory_path)
        
        if not path.is_dir():
            raise ValueError(f"Not a directory: {directory_path}")
        
        all_documents = []
        pattern = "**/*" if recursive else "*"
        
        for file_path in path.glob(pattern):
            if file_path.is_file():
                extension = file_path.suffix.lower().lstrip('.')
                
                if extension in self.SUPPORTED_EXTENSIONS:
                    try:
                        print(f"Loading: {file_path}")
                        docs = self.load_document(str(file_path))
                        all_documents.extend(docs)
                        print(f"✓ Loaded {len(docs)} document(s) from {file_path.name}")
                    except Exception as e:
                        print(f"✗ Failed to load {file_path.name}: {e}")
        
        return all_documents
    
    @classmethod
    def get_supported_extensions(cls) -> List[str]:
        """Get list of supported file extensions"""
        return list(cls.SUPPORTED_EXTENSIONS.keys())


if __name__ == "__main__":
    # Test the document loader
    loader = DocumentLoader()
    print(f"Supported extensions: {loader.get_supported_extensions()}")
