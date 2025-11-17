"""Excel logging helpers for the Streamlit agent."""

from __future__ import annotations

import os
import re
from datetime import datetime
from difflib import SequenceMatcher
from typing import Any, Iterable, List, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .llm_cache import cached_tool


EXCEL_LOG_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "logs", "agent_logs.xlsx")


@cached_tool(ttl_seconds=60 * 30, cache_type="excel_get_cached_response", write_to_query_store=True)
def get_cached_response(query: str, threshold: float = 0.8) -> Optional[dict]:
    """Return a cached response whose query is similar to *query*.

    Args:
        query: User text to find in the log.
        threshold: Minimum similarity (0-1) required to reuse a response.

    Returns:
        Matching log entry metadata or ``None`` when no suitable match exists.
    """

    print(f"\n[get_cached_response] Input: query={query}, threshold={threshold}")
    if not os.path.exists(EXCEL_LOG_PATH):
        print("[get_cached_response] Excel log not found.")
        return None

    wb = load_workbook(EXCEL_LOG_PATH)
    ws = wb.active

    best_match = None
    best_score = 0.0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        user_input = row[1]
        score = SequenceMatcher(None, query, user_input).ratio()
        if score > best_score and score >= threshold:
            best_score = score
            best_match = {
                "timestamp": row[0],
                "user_input": user_input,
                "output": row[4],
                "status": row[5],
                "similarity": score,
            }

    print(f"[get_cached_response] Output: {best_match}")
    return best_match


@cached_tool(ttl_seconds=60 * 5, cache_type="excel_history", write_to_query_store=True)
def get_conversation_history(last_n: int = 5, all_history: bool = False) -> List[dict]:
    """Return recent conversation history from the Excel log.

    Args:
        last_n: Maximum number of entries to return.
        all_history: When ``True`` return every stored interaction.

    Returns:
        List of dictionaries with ``timestamp``, ``user_input``, ``tools_used``,
        ``sources``, ``output``, and ``status`` keys.
    """

    print(f"\n[get_conversation_history] Input: last_n={last_n}, all_history={all_history}")
    if not os.path.exists(EXCEL_LOG_PATH):
        print("[get_conversation_history] Excel log not found.")
        return []

    wb = load_workbook(EXCEL_LOG_PATH)
    ws = wb.active

    history = [
        {
            "timestamp": row[0],
            "user_input": row[1],
            "tools_used": row[2],
            "sources": row[3],
            "output": row[4],
            "status": row[5],
        }
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True)
    ]

    result = history if all_history else history[-last_n:]
    print(f"[get_conversation_history] Output: {result}")
    return history if all_history else result


def initialize_excel_log() -> None:
    """Create the Excel log with headers when it does not already exist."""

    print("\n[initialize_excel_log] Called")
    if os.path.exists(EXCEL_LOG_PATH):
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Agent Logs"

    headers = ["Timestamp", "User Input", "Tools Used", "Sources/URLs", "Output", "Status"]
    ws.append(headers)

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["F"].width = 12

    wb.save(EXCEL_LOG_PATH)
    print(f"Excel log created at: {EXCEL_LOG_PATH}")


def extract_urls(text: str) -> List[str]:
    """Extract hyperlinks from an arbitrary chunk of text."""

    url_pattern = r"https?://[^\s<>\"{}|\\^`\[\]]+"
    return re.findall(url_pattern, text)


def _unique(iterable: Iterable[str]) -> List[str]:
    seen = set()
    result = []
    for item in iterable:
        if item not in seen:
            seen.add(item)
            result.append(item)
    return result


def log_interaction(
    user_input: str,
    tools_used: List[str],
    agent_messages: List[Any],
    output: str,
    status: str = "Success",
) -> bool:
    """Append an interaction row to the Excel log.

    Args:
        user_input: Prompt text provided by the user.
        tools_used: Names of tools invoked while handling the request.
        agent_messages: Full message history (used to extract cited URLs).
        output: Final response text returned to the user.
        status: ``"Success"`` or ``"Error"`` indicator.

    Returns:
        ``True`` when the row was written successfully, otherwise ``False``.
    """

    try:
        initialize_excel_log()

        wb = load_workbook(EXCEL_LOG_PATH)
        ws = wb.active

        sources: List[str] = []
        for msg in agent_messages:
            if hasattr(msg, "content") and msg.content:
                sources.extend(extract_urls(str(msg.content)))
        sources = _unique(sources)

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        row_data = [
            timestamp,
            user_input[:500],
            ", ".join(tools_used) if tools_used else "None",
            "\n".join(sources[:5]) if sources else "No URLs",
            output[:1000] if output else "No output",
            status,
        ]

        ws.append(row_data)

        row_num = ws.max_row
        for cell in ws[row_num]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        status_cell = ws.cell(row=row_num, column=6)
        if status == "Success":
            status_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            status_cell.font = Font(color="006100")
        else:
            status_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            status_cell.font = Font(color="9C0006")

        wb.save(EXCEL_LOG_PATH)
        print(f"Logged to Excel: {EXCEL_LOG_PATH}")
        return True

    except Exception as exc:  # pragma: no cover - defensive logging
        print(f"Error logging to Excel: {exc}")
        return False


def get_log_summary() -> str:
    """Return human-readable summary statistics for the Excel log."""

    try:
        if not os.path.exists(EXCEL_LOG_PATH):
            return "No log file exists yet"

        wb = load_workbook(EXCEL_LOG_PATH)
        ws = wb.active

        total_queries = ws.max_row - 1
        if total_queries <= 0:
            return "No queries logged yet"

        success_count = 0
        error_count = 0
        for row in ws.iter_rows(min_row=2, max_col=6, max_row=ws.max_row):
            status_value = row[5].value
            if status_value == "Success":
                success_count += 1
            else:
                error_count += 1

        summary = (
            "Log Summary:\n"
            f"- Total Queries: {total_queries}\n"
            f"- Successful: {success_count}\n"
            f"- Errors: {error_count}\n"
            f"- Log File: {EXCEL_LOG_PATH}"
        )
        return summary

    except Exception as exc:  # pragma: no cover - defensive logging
        return f"Error reading log: {exc}"
